"""Word layouts found in real drafts: tables after the reference list, and mixed heading styles."""

from pathlib import Path

import docx

from paper_check.ingest import ingest
from paper_check.model import PaperMap


def _read(path: Path, tmp_path: Path) -> PaperMap:
    ingest(path, output=str(tmp_path / "out"))
    return PaperMap.load(tmp_path / "out" / "paper-check-work" / "paper_map.json")


def test_tables_after_the_reference_list_leave_the_references(tmp_path):
    d = docx.Document()
    d.add_heading("A Paper", 0)
    d.add_heading("1. Introduction", 1)
    d.add_paragraph("We study savings (Smith, 2020; Jones, 2019).")
    d.add_heading("References", 1)
    d.add_paragraph("Jones, A. (2019). Saving more. Journal of Savings, 3(1), 1-10.")
    d.add_paragraph("Smith, B. (2020). Saving less. Journal of Savings, 4(2), 11-20.")
    d.add_paragraph("Table 1 Summary Statistics for the Primary Sample")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Variable", "Mean"
    t.cell(1, 0).text, t.cell(1, 1).text = "Income", "52.1"
    d.add_paragraph("Note: Means are in thousands of dollars.")
    path = tmp_path / "paper.docx"
    d.save(path)
    pm = _read(path, tmp_path)
    assert len(pm.references) == 2
    table = [b for b in pm.blocks if b.kind == "table"][0]
    caption = [b for b in pm.blocks if b.kind == "caption"][0]
    assert table.section == ["Tables and figures"] and caption.text.startswith("Table 1")
    assert not any(b.kind == "reference" and "Summary Statistics" in b.text for b in pm.blocks)


def test_numbered_headings_take_their_level_from_the_number(tmp_path):
    d = docx.Document()
    d.add_heading("A Paper", 0)
    d.add_heading("4. Models", 1)
    d.add_heading("4.1 First analysis", 2)
    d.add_paragraph("Text of the first analysis.")
    d.add_heading("4.2 Second analysis", 3)  # wrong style level in the draft
    d.add_paragraph("Text of the second analysis.")
    path = tmp_path / "paper.docx"
    d.save(path)
    pm = _read(path, tmp_path)
    second = [b for b in pm.blocks if b.kind == "paragraph" and b.text.startswith("Text of the second")][0]
    assert second.section == ["4. Models", "4.2 Second analysis"]


def test_a_typed_appendix_table_title_is_its_caption(tmp_path):
    from paper_check.checks import Context, run_checks

    d = docx.Document()
    d.add_heading("A Paper", 0)
    d.add_heading("1. Results", 1)
    d.add_paragraph("Savings rose (both p < 0.01; Appendix Table A1). Table 7 has more.")
    d.add_heading("Appendix", 1)
    d.add_paragraph("Appendix Table A1 Robustness of the Main Estimates")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Variable", "Estimate"
    t.cell(1, 0).text, t.cell(1, 1).text = "Share", "0.12"
    path = tmp_path / "paper.docx"
    d.save(path)
    pm = _read(path, tmp_path)
    caption = [b for b in pm.blocks if b.kind == "caption"]
    assert caption and caption[0].text.startswith("Appendix Table A1")
    mention = [b for b in pm.blocks if b.text.startswith("Savings rose")][0]
    refs = {r.raw: r.resolved for r in mention.crossrefs}
    assert refs.get("Appendix Table A1") is True and refs.get("Table 7") is False
    work = tmp_path / "out" / "paper-check-work"
    found = run_checks(pm, Context(work=work)).findings
    fmt05 = [f for f in found if f.check_id == "FMT-05"]
    assert [f.quote for f in fmt05 if "A1" in (f.quote or "")] == []
    table7 = [f for f in fmt05 if "Table 7" in (f.quote or "")]
    assert table7 and "LaTeX" not in table7[0].problem and "no caption, title or heading" in table7[0].problem
    assert not [f for f in found if f.check_id in ("FMT-06", "FMT-07") and "A1" in (f.quote or "")]


def test_a_references_title_in_the_bibliography_style_is_a_heading(tmp_path):
    d = docx.Document()
    try:
        d.styles.add_style("Bibliography", 1)
    except ValueError:
        pass
    d.add_heading("A Paper", 0)
    d.add_heading("1. Introduction", 1)
    d.add_paragraph("We study savings (Smith, 2020; Jones, 2019).")
    d.add_paragraph("References:", style="Bibliography")
    d.add_paragraph("Jones, A. (2019). Saving more. Journal of Savings, 3(1), 1-10.", style="Bibliography")
    d.add_paragraph("Smith, B. (2020). Saving less. Journal of Savings, 4(2), 11-20.", style="Bibliography")
    path = tmp_path / "paper.docx"
    d.save(path)
    pm = _read(path, tmp_path)
    assert [r.raw[:5] for r in pm.references] == ["Jones", "Smith"]
    assert any(b.kind == "heading" and b.text == "References:" for b in pm.blocks)
